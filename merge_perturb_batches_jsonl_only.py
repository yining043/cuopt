#!/usr/bin/env python3
"""
Lightweight merge for perturbation batches (JSONL + Excel).
Uses openpyxl directly (no pandas) to avoid segfault in certain environments.

JSONL dedup key: (initial_solution.edges_hash, perturbation_step)
Excel dedup key: original_edges_hash
Keep strategy: last wins
"""

import argparse
import glob
import json
import os
import tempfile
from typing import Dict, Tuple, List, Optional


def _iter_jsonl(path: str):
    with open(path, "r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if line:
                try:
                    yield json.loads(line)
                except Exception:
                    pass


def _record_key(rec: dict) -> Optional[Tuple[str, int]]:
    init = rec.get("initial_solution") or {}
    h = init.get("edges_hash")
    if not h:
        return None
    step = int(rec.get("perturbation_step") or 0)
    return (str(h), step) if step > 0 else None


def merge_jsonl(out_dir: str, operator_type: str, runs: int, batch_suffix: str = None):
    all_jsonl = os.path.join(out_dir, f"{operator_type}_training_data.ALL_r{runs}.jsonl")
    # Support custom batch suffix pattern, or match both _r30 and _r30_r30
    if batch_suffix:
        batch_pattern = os.path.join(out_dir, f"{operator_type}_training_data.batch_*{batch_suffix}.jsonl")
        batch_files = sorted(glob.glob(batch_pattern))
    else:
        batch_pattern = os.path.join(out_dir, f"{operator_type}_training_data.batch_*_r{runs}*.jsonl")
        batch_files = sorted(glob.glob(batch_pattern))

    inputs: List[str] = []
    if os.path.exists(all_jsonl):
        inputs.append(all_jsonl)
    inputs.extend(batch_files)

    if not inputs:
        os.makedirs(out_dir, exist_ok=True)
        open(all_jsonl, "a", encoding="utf-8").close()
        print(f"[JSONL] No files to merge, created empty: {all_jsonl}")
        return all_jsonl, []

    # Dedup: key -> record
    merged: Dict[Tuple[str, int], dict] = {}
    for path in inputs:
        for rec in _iter_jsonl(path):
            k = _record_key(rec)
            if k:
                merged[k] = rec  # last wins

    # Stable order by (edges_hash, step)
    items = sorted(merged.items(), key=lambda kv: (kv[0][0], kv[0][1]))

    # Atomic write
    os.makedirs(out_dir, exist_ok=True)
    tmp_fd, tmp_path = tempfile.mkstemp(prefix=f".{operator_type}_ALL_", suffix=".jsonl", dir=out_dir)
    try:
        with os.fdopen(tmp_fd, "w", encoding="utf-8") as w:
            for _, rec in items:
                w.write(json.dumps(rec, ensure_ascii=False) + "\n")
        os.replace(tmp_path, all_jsonl)
    except Exception:
        try:
            os.remove(tmp_path)
        except Exception:
            pass
        raise

    print(f"[JSONL] Merged {len(inputs)} files -> {all_jsonl} ({len(merged)} unique records)")
    return all_jsonl, batch_files


def merge_xlsx(out_dir: str, operator_type: str, runs: int, batch_suffix: str = None):
    """Merge Excel files using openpyxl directly (no pandas)."""
    from openpyxl import Workbook, load_workbook

    all_xlsx = os.path.join(out_dir, f"{operator_type}_results.ALL_r{runs}.xlsx")
    if batch_suffix:
        batch_pattern = os.path.join(out_dir, f"{operator_type}_results.batch_*{batch_suffix}.xlsx")
        batch_files = sorted(glob.glob(batch_pattern))
    else:
        batch_pattern = os.path.join(out_dir, f"{operator_type}_results.batch_*_r{runs}*.xlsx")
        batch_files = sorted(glob.glob(batch_pattern))

    inputs: List[str] = []
    if os.path.exists(all_xlsx):
        inputs.append(all_xlsx)
    inputs.extend(batch_files)

    if not inputs:
        os.makedirs(out_dir, exist_ok=True)
        wb = Workbook()
        wb.save(all_xlsx)
        print(f"[XLSX] No files to merge, created empty: {all_xlsx}")
        return all_xlsx, []

    # Read all rows, dedup by original_edges_hash
    header = None
    hash_col_idx = None
    merged: Dict[str, List] = {}  # edges_hash -> row values

    for path in inputs:
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            wb.close()
        except Exception as e:
            print(f"[XLSX] Warning: failed to read {path}: {e}")
            continue

        if not rows:
            continue

        # First file sets the header
        if header is None:
            header = list(rows[0])
            try:
                hash_col_idx = header.index("original_edges_hash")
            except ValueError:
                hash_col_idx = None
            data_rows = rows[1:]
        else:
            # Skip header row if it matches
            if rows[0] == tuple(header):
                data_rows = rows[1:]
            else:
                data_rows = rows

        for row in data_rows:
            if hash_col_idx is not None and len(row) > hash_col_idx:
                key = str(row[hash_col_idx]) if row[hash_col_idx] else ""
            else:
                key = str(row)  # fallback: entire row as key
            if key:
                merged[key] = list(row)  # last wins

    # Write merged Excel
    os.makedirs(out_dir, exist_ok=True)
    wb_out = Workbook()
    ws_out = wb_out.active
    ws_out.title = "Perturbation Results"

    if header:
        ws_out.append(header)
    for row in merged.values():
        ws_out.append(row)

    tmp_path = os.path.join(out_dir, f".{operator_type}_results.ALL_r{runs}.tmp.xlsx")
    try:
        wb_out.save(tmp_path)
        os.replace(tmp_path, all_xlsx)
    except Exception:
        try:
            os.remove(tmp_path)
        except Exception:
            pass
        raise

    print(f"[XLSX] Merged {len(inputs)} files -> {all_xlsx} ({len(merged)} unique rows)")
    return all_xlsx, batch_files


def delete_batch_files(jsonl_batches: List[str], xlsx_batches: List[str], dry_run: bool = False):
    for f in jsonl_batches + xlsx_batches:
        if dry_run:
            print(f"  [dry-run] would delete: {f}")
        else:
            try:
                os.remove(f)
                print(f"  Deleted: {f}")
            except Exception as e:
                print(f"  Failed to delete {f}: {e}")


def rebuild_xlsx_from_jsonl(out_dir: str, operator_type: str, runs: int):
    """Rebuild Excel from JSONL (ensures all data is included)."""
    from openpyxl import Workbook
    
    all_jsonl = os.path.join(out_dir, f"{operator_type}_training_data.ALL_r{runs}.jsonl")
    all_xlsx = os.path.join(out_dir, f"{operator_type}_results.ALL_r{runs}.xlsx")
    
    if not os.path.exists(all_jsonl):
        print(f"[XLSX] JSONL not found: {all_jsonl}")
        return None
    
    # Aggregate by edges_hash (one row per optima)
    optima: Dict[str, dict] = {}
    for rec in _iter_jsonl(all_jsonl):
        init = rec.get("initial_solution") or {}
        h = init.get("edges_hash")
        if not h:
            continue
        
        step = rec.get("perturbation_step", 0)
        ps = rec.get("perturbed_solution") or {}
        
        if h not in optima:
            optima[h] = {
                "original_edges_hash": h,
                "original_cost": init.get("cost"),
                "k": rec.get("k", 5),
            }
        
        # Add step-specific data
        prefix = f"step{step}_"
        optima[h][f"{prefix}perturbed_cost"] = ps.get("cost")
        optima[h][f"{prefix}return_to_original_ratio"] = rec.get("return_to_original_ratio")
        optima[h][f"{prefix}jaccard_distance"] = ps.get("jaccard_distance")
        optima[h][f"{prefix}broken_pairs_ratio"] = ps.get("broken_pairs_ratio")
    
    if not optima:
        print(f"[XLSX] No data in JSONL")
        return None
    
    # Build header
    sample = list(optima.values())[0]
    header = ["original_edges_hash", "original_cost", "k"]
    for step in range(1, 6):
        prefix = f"step{step}_"
        header.extend([
            f"{prefix}perturbed_cost",
            f"{prefix}return_to_original_ratio", 
            f"{prefix}jaccard_distance",
            f"{prefix}broken_pairs_ratio",
        ])
    
    # Write Excel
    wb = Workbook()
    ws = wb.active
    ws.title = "Perturbation Results"
    ws.append(header)
    
    for h in sorted(optima.keys()):
        row = [optima[h].get(col) for col in header]
        ws.append(row)
    
    tmp_path = all_xlsx + ".tmp"
    wb.save(tmp_path)
    os.replace(tmp_path, all_xlsx)
    
    print(f"[XLSX] Rebuilt from JSONL: {all_xlsx} ({len(optima)} rows)")
    return all_xlsx


def main():
    parser = argparse.ArgumentParser(description="Merge perturbation batch files (JSONL + Excel, no pandas)")
    parser.add_argument("--out_dir", required=True, help="Output directory")
    parser.add_argument("--operator_type", default="remove_and_insert")
    parser.add_argument("--runs", type=int, default=30)
    parser.add_argument("--batch_suffix", default=None, help="Custom batch file suffix (e.g., '_r30_r30')")
    parser.add_argument("--delete_batches", action="store_true", help="Delete batch files after merge")
    parser.add_argument("--dry_run", action="store_true", help="Show what would be deleted without deleting")
    parser.add_argument("--rebuild_xlsx", action="store_true", help="Rebuild Excel from JSONL instead of merging xlsx files")
    args = parser.parse_args()

    _, jsonl_batches = merge_jsonl(args.out_dir, args.operator_type, args.runs, args.batch_suffix)
    
    if args.rebuild_xlsx:
        rebuild_xlsx_from_jsonl(args.out_dir, args.operator_type, args.runs)
        # Still collect xlsx batch files for deletion (both suffix patterns)
        if args.batch_suffix:
            xlsx_pattern = os.path.join(args.out_dir, f"{args.operator_type}_results.batch_*{args.batch_suffix}.xlsx")
        else:
            xlsx_pattern = os.path.join(args.out_dir, f"{args.operator_type}_results.batch_*_r{args.runs}*.xlsx")
        xlsx_batches = sorted(glob.glob(xlsx_pattern))
    else:
        _, xlsx_batches = merge_xlsx(args.out_dir, args.operator_type, args.runs, args.batch_suffix)

    if (jsonl_batches or xlsx_batches) and (args.delete_batches or args.dry_run):
        delete_batch_files(jsonl_batches, xlsx_batches, dry_run=args.dry_run)


if __name__ == "__main__":
    main()
