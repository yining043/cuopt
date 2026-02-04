#!/usr/bin/env python3
"""
Analyze data from the first 10 runs:
Each record represents an initial solution run 100 times, reaching different basins
For each record, remove basins with count < 10
Generate Excel table and heatmap
"""

import json
import os
from collections import defaultdict, Counter
import matplotlib.pyplot as plt
import seaborn as sns
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


def read_training_data(file_path, max_runs=10):
    """
    Read training_data.jsonl file, keep only data from the first max_runs runs
    """
    data = []
    seen_runs = set()
    
    with open(file_path, 'r', encoding='utf-8') as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            
            try:
                rec = json.loads(line)
                run_id = rec.get('run_id')
                
                # Keep only the first max_runs different runs
                if run_id not in seen_runs:
                    if len(seen_runs) >= max_runs:
                        break
                    seen_runs.add(run_id)
                
                # If this run_id is in the first max_runs, keep it
                if run_id in seen_runs:
                    data.append(rec)
            except Exception as e:
                print(f"Error parsing line: {e}")
                continue
    
    return data, seen_runs


def process_single_record(rec, min_count=10):
    """
    Process a single record:
    1. Convert probabilities in basin_distribution to counts (probability * 100)
    2. Filter out basins with count < min_count
    3. Sort by count
    Returns cleaned basin information
    """
    basin_dist = rec.get('basin_distribution', {})
    basin_features = rec.get('basin_features', {})
    
    if not basin_dist:
        return None
    
    # Convert probabilities to counts (assuming 100 total runs)
    basin_counts = {}
    for basin_hash, prob in basin_dist.items():
        count = int(round(prob * 100))  # Convert probability to count
        if count >= min_count:
            basin_counts[basin_hash] = count
    
    if not basin_counts:
        return None
    
    # Sort by count
    sorted_basins = sorted(basin_counts.items(), key=lambda x: x[1], reverse=True)
    
    # Collect detailed information for each basin
    basin_info = []
    for basin_hash, count in sorted_basins:
        features = basin_features.get(basin_hash, {})
        basin_info.append({
            'basin_hash': basin_hash,
            'count': count,
            'cost': features.get('mean_cost', 'N/A'),
            'gap_to_hgs': features.get('gap_to_hgs', 'N/A')
        })
    
    return {
        'basin_count': len(basin_info),
        'basins': basin_info
    }


def create_excel(data, processed_records, output_path):
    """
    Create Excel table
    Each row contains: run_id, trial_id, global_iter, local_iter, original cost, hash,
    cleaned basin count, and for each basin: hash, cost, count
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Basin Analysis"
    
    # Headers
    headers = ['run_id', 'trial_id', 'global_iter', 'local_iter', 
               'initial_cost', 'initial_hash', 'cleaned_basin_count']
    
    # Find maximum number of basins to determine how many columns are needed
    max_basins = max([pr['basin_count'] for pr in processed_records.values() if pr is not None] + [0])
    
    # Add 3 columns for each basin: hash, cost, count
    for i in range(max_basins):
        headers.extend([f'basin_{i+1}_hash', f'basin_{i+1}_cost', f'basin_{i+1}_count'])
    
    # Write headers
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center')
    
    # Write data
    row_idx = 2
    for i, rec in enumerate(data):
        if i not in processed_records:
            continue
        
        pr = processed_records[i]
        if pr is None:
            continue
        
        basins = pr['basins']
        
        # Basic information
        initial_solution = rec.get('initial_solution', {})
        ws.cell(row=row_idx, column=1, value=rec.get('run_id', ''))
        ws.cell(row=row_idx, column=2, value=rec.get('trial_id', ''))
        ws.cell(row=row_idx, column=3, value=rec.get('global_iter', ''))
        ws.cell(row=row_idx, column=4, value=rec.get('local_iter', ''))
        ws.cell(row=row_idx, column=5, value=initial_solution.get('cost', ''))
        ws.cell(row=row_idx, column=6, value=initial_solution.get('edges_hash', ''))
        ws.cell(row=row_idx, column=7, value=pr['basin_count'])
        
        # Basin information
        col_idx = 8
        for basin in basins:
            ws.cell(row=row_idx, column=col_idx, value=basin['basin_hash'])
            col_idx += 1
            ws.cell(row=row_idx, column=col_idx, value=basin['cost'])
            col_idx += 1
            ws.cell(row=row_idx, column=col_idx, value=basin['count'])
            col_idx += 1
        
        row_idx += 1
    
    # Adjust column width
    for col_idx in range(1, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        if col_idx <= 7:
            ws.column_dimensions[col_letter].width = 15
        else:
            ws.column_dimensions[col_letter].width = 20
    
    wb.save(output_path)
    print(f"Excel file saved to: {output_path}")


def create_heatmap_data(processed_records, data):
    """
    Create heatmap data
    X-axis: basin hash
    Y-axis: data records (run_id + trial_id)
    """
    # Collect all basins that appear (in cleaned data)
    # Only keep records that fall into more than one basin
    all_basins = set()
    valid_records = []
    
    for i, pr in processed_records.items():
        if pr and pr['basin_count'] > 1:  # Only keep records with more than one basin
            valid_records.append(i)
            for basin in pr['basins']:
                all_basins.add(basin['basin_hash'])
    
    if len(all_basins) <= 2:
        print(f"Warning: Number of cleaned basins ({len(all_basins)}) <= 2, skipping heatmap")
        return None, None, None
    
    # Sort basin hashes
    sorted_basins = sorted(all_basins)
    
    # Create matrix
    n_records = len(valid_records)
    n_basins = len(sorted_basins)
    
    heatmap_matrix = np.zeros((n_records, n_basins))
    row_labels = []
    
    for idx, rec_idx in enumerate(valid_records):
        rec = data[rec_idx]
        pr = processed_records[rec_idx]
        
        # Row label
        row_label = f"{rec.get('run_id', '')}_t{rec.get('trial_id', '')}"
        row_labels.append(row_label)
        
        # Fill in counts (can be normalized to 0-1 range for display)
        for basin in pr['basins']:
            basin_hash = basin['basin_hash']
            if basin_hash in sorted_basins:
                j = sorted_basins.index(basin_hash)
                # Use count as value (can normalize, but using count directly here)
                heatmap_matrix[idx, j] = basin['count']
    
    return heatmap_matrix, sorted_basins, row_labels


def create_top_basins_heatmap_data(processed_records, data, top_n=6):
    """
    Create heatmap data for top N basins
    X-axis: top 1, top 2, ..., top N basin positions
    Y-axis: data records (run_id + trial_id)
    Value: count of the basin at that position
    """
    # Only keep records that fall into more than one basin
    valid_records = []
    
    for i, pr in processed_records.items():
        if pr and pr['basin_count'] > 1:  # Only keep records with more than one basin
            valid_records.append(i)
    
    if len(valid_records) == 0:
        print(f"Warning: No records with multiple basins, skipping top basins heatmap")
        return None, None, None
    
    # Create matrix: rows = records, cols = top positions (1 to top_n)
    n_records = len(valid_records)
    heatmap_matrix = np.zeros((n_records, top_n))
    row_labels = []
    
    for idx, rec_idx in enumerate(valid_records):
        rec = data[rec_idx]
        pr = processed_records[rec_idx]
        
        # Row label
        row_label = f"{rec.get('run_id', '')}_t{rec.get('trial_id', '')}"
        row_labels.append(row_label)
        
        # Fill in counts for top N basins (basins are already sorted by count)
        for pos in range(min(top_n, len(pr['basins']))):
            heatmap_matrix[idx, pos] = pr['basins'][pos]['count']
    
    # Column labels
    col_labels = [f'Top {i+1}' for i in range(top_n)]
    
    return heatmap_matrix, col_labels, row_labels


def plot_heatmap(heatmap_matrix, basin_hashes, row_labels, output_path, max_rows=5000):
    """
    Plot heatmap
    """
    n_rows, n_cols = heatmap_matrix.shape
    
    # If too many rows, sample
    if n_rows > max_rows:
        print(f"  Number of rows ({n_rows}) exceeds limit ({max_rows}), sampling...")
        step = n_rows // max_rows
        indices = np.arange(0, n_rows, step)[:max_rows]
        heatmap_matrix = heatmap_matrix[indices]
        row_labels = [row_labels[i] for i in indices]
        n_rows = len(row_labels)
        print(f"  Number of rows after sampling: {n_rows}")
    
    # If too many basins, keep only the first 100
    if n_cols > 100:
        print(f"  Number of basins ({n_cols}) exceeds limit (100), keeping only first 100...")
        heatmap_matrix = heatmap_matrix[:, :100]
        basin_hashes = basin_hashes[:100]
        n_cols = 100
    
    # Truncate basin hash for display
    display_basin_hashes = [h[:8] + '...' if len(h) > 8 else h for h in basin_hashes]
    
    # Simplify row labels
    if len(row_labels) > 100:
        display_row_labels = [f"R{i}" for i in range(len(row_labels))]
    else:
        display_row_labels = row_labels
    
    # Set figure size
    fig_width = min(20, max(12, n_cols * 0.3))
    fig_height = min(30, max(8, n_rows * 0.1))
    
    # Create DataFrame
    df = pd.DataFrame(heatmap_matrix, 
                      index=display_row_labels,
                      columns=display_basin_hashes)
    
    fig, ax = plt.subplots(figsize=(fig_width, fig_height))
    
    # Use seaborn to plot heatmap
    sns.heatmap(df, 
                cmap='YlOrRd', 
                annot=False,
                fmt='.0f',
                cbar_kws={'label': 'Count'},
                xticklabels=True if n_cols <= 50 else False,
                yticklabels=True if n_rows <= 100 else False,
                ax=ax)
    
    plt.title(f'Basin Distribution Heatmap (Cleaned Data)\n({n_rows} records, {n_cols} basins)', fontsize=12)
    plt.xlabel('Basin Hash (truncated)', fontsize=10)
    plt.ylabel('Data Records', fontsize=10)
    plt.tight_layout()
    
    plt.savefig(output_path, dpi=150, bbox_inches='tight')
    print(f"Heatmap saved to: {output_path}")
    plt.close()


def main():
    file_path = "/home/jieyi/cuopt/basin_datasets0_analyze/cvrp100_uniform.pkl#0/training_data.jsonl"
    max_runs = 10
    min_count = 10
    
    print(f"Reading file: {file_path}")
    print(f"Keeping only data from the first {max_runs} runs")
    print(f"Filtering basins with count < {min_count}")
    
    # Read data
    data, seen_runs = read_training_data(file_path, max_runs=max_runs)
    print(f"Read {len(data)} records from {len(seen_runs)} runs")
    print(f"Runs: {sorted(seen_runs)}")
    
    # Process each record
    print(f"\nProcessing each record, filtering basins with count < {min_count}...")
    processed_records = {}
    valid_count = 0
    
    for i, rec in enumerate(data):
        pr = process_single_record(rec, min_count=min_count)
        processed_records[i] = pr
        if pr and pr['basin_count'] > 0:
            valid_count += 1
    
    print(f"Valid records (with basins after cleaning): {valid_count}/{len(data)}")
    
    # Statistics of cleaned basins
    all_cleaned_basins = set()
    total_basins_before = 0
    total_basins_after = 0
    records_with_single_basin = 0
    records_with_multiple_basins = 0
    
    for pr in processed_records.values():
        if pr:
            total_basins_before += len(pr.get('basins', []))  # Should be before cleaning, but we didn't save it
            total_basins_after += pr['basin_count']
            if pr['basin_count'] == 1:
                records_with_single_basin += 1
            elif pr['basin_count'] > 1:
                records_with_multiple_basins += 1
            for basin in pr['basins']:
                all_cleaned_basins.add(basin['basin_hash'])
    
    print(f"Total number of different basins after cleaning: {len(all_cleaned_basins)}")
    print(f"Total number of basins in all records after cleaning: {total_basins_after}")
    print(f"Records with single basin (will be excluded from heatmap): {records_with_single_basin}")
    print(f"Records with multiple basins (will be included in heatmap): {records_with_multiple_basins}")
    
    # Generate Excel
    excel_path = "/home/jieyi/cuopt/basin_analysis.xlsx"
    print(f"\nGenerating Excel table...")
    create_excel(data, processed_records, excel_path)
    
    # Create heatmap data (excluding records that fall into only one basin)
    print(f"\nCreating heatmap data (excluding records with only one basin)...")
    heatmap_matrix, sorted_basins, row_labels = create_heatmap_data(processed_records, data)
    
    if heatmap_matrix is not None:
        print(f"Heatmap matrix size: {heatmap_matrix.shape}")
        print(f"Records in heatmap (with >1 basin): {len(row_labels)}")
        
        # Plot heatmap
        heatmap_path = "/home/jieyi/cuopt/basin_heatmap.png"
        print(f"\nPlotting heatmap...")
        plot_heatmap(heatmap_matrix, sorted_basins, row_labels, heatmap_path)
    else:
        print("Skipping heatmap generation")
    
    # Create top basins heatmap
    print(f"\nCreating top basins heatmap data...")
    top_heatmap_matrix, top_col_labels, top_row_labels = create_top_basins_heatmap_data(processed_records, data, top_n=6)
    
    if top_heatmap_matrix is not None:
        print(f"Top basins heatmap matrix size: {top_heatmap_matrix.shape}")
        print(f"Records in top basins heatmap: {len(top_row_labels)}")
        
        # Plot top basins heatmap
        top_heatmap_path = "/home/jieyi/cuopt/basin_heatmap_top6.png"
        print(f"\nPlotting top basins heatmap...")
        plot_heatmap(top_heatmap_matrix, top_col_labels, top_row_labels, top_heatmap_path, max_rows=5000)
    else:
        print("Skipping top basins heatmap generation")
    
    print("\nDone!")


if __name__ == "__main__":
    main()
